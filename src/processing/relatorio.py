import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os

from src.utils.logger import configurar_logger

logger = configurar_logger("relatorio")


def gerar_relatorio(df_validos: pd.DataFrame, df_invalidos: pd.DataFrame, caminho_saida: str):
    """
    Gera um relatório Excel formatado com:
    - Aba 'Resumo' com contagens gerais
    - Aba 'Válidos' com os registros aprovados
    - Aba 'Inválidos' com os registros com erro e o motivo
    """
    os.makedirs(os.path.dirname(caminho_saida), exist_ok=True)

    wb = Workbook()

    # --- Aba Resumo ---
    ws_resumo = wb.active
    ws_resumo.title = "Resumo"

    titulo_font = Font(bold=True, size=14)
    cabecalho_font = Font(bold=True, color="FFFFFF")
    cabecalho_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    ws_resumo["A1"] = "Relatório de Processamento de Solicitações"
    ws_resumo["A1"].font = titulo_font
    ws_resumo.merge_cells("A1:B1")

    total = len(df_validos) + len(df_invalidos)
    linhas_resumo = [
        ("Total de solicitações", total),
        ("Aprovadas (válidas)", len(df_validos)),
        ("Com erro (inválidas)", len(df_invalidos)),
    ]

    ws_resumo["A3"] = "Indicador"
    ws_resumo["B3"] = "Quantidade"
    ws_resumo["A3"].font = cabecalho_font
    ws_resumo["B3"].font = cabecalho_font
    ws_resumo["A3"].fill = cabecalho_fill
    ws_resumo["B3"].fill = cabecalho_fill

    for i, (label, valor) in enumerate(linhas_resumo, start=4):
        ws_resumo[f"A{i}"] = label
        ws_resumo[f"B{i}"] = valor

    ws_resumo.column_dimensions["A"].width = 28
    ws_resumo.column_dimensions["B"].width = 15

    # --- Aba Válidos ---
    ws_validos = wb.create_sheet("Válidos")
    _escrever_dataframe(ws_validos, df_validos, cabecalho_font, cabecalho_fill)

    # --- Aba Inválidos ---
    ws_invalidos = wb.create_sheet("Inválidos")
    _escrever_dataframe(ws_invalidos, df_invalidos, cabecalho_font, cabecalho_fill)

    try:
        wb.save(caminho_saida)
        logger.info(f"Relatório salvo em: {caminho_saida}")
    except Exception as erro:
        logger.error(f"Erro ao salvar relatório em {caminho_saida}: {erro}")
        raise


def _escrever_dataframe(ws, df: pd.DataFrame, cabecalho_font, cabecalho_fill):
    """Escreve um DataFrame numa planilha do openpyxl, com cabeçalho formatado."""
    if df.empty:
        ws["A1"] = "Nenhum registro"
        return

    for col_idx, coluna in enumerate(df.columns, start=1):
        celula = ws.cell(row=1, column=col_idx, value=coluna)
        celula.font = cabecalho_font
        celula.fill = cabecalho_fill

    for row_idx, linha in enumerate(df.itertuples(index=False), start=2):
        for col_idx, valor in enumerate(linha, start=1):
            ws.cell(row=row_idx, column=col_idx, value=valor)

    for col_idx, coluna in enumerate(df.columns, start=1):
        largura = max(15, len(str(coluna)) + 2)
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = largura


# Teste rápido
if __name__ == "__main__":
    from src.processing.validador import carregar_e_validar

    validos, invalidos = carregar_e_validar("data/input/solicitacoes.xlsx")
    gerar_relatorio(validos, invalidos, "data/output/relatorio_final.xlsx")